from tkinter import *
from tkinter import messagebox,Toplevel,filedialog
from tkinter import ttk
from tkcalendar import DateEntry
from datetime import datetime
import pyodbc
from openpyxl import Workbook
from docx import Document
import os
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt
import webbrowser

#===== HÀM KẾT NỐI SQL =====
try:
    conn = pyodbc.connect(
        'DRIVER={SQL Server};'
        'SERVER=DESKTOP-MA5J22U\\SQLSERVER2022;'
        'DATABASE=QLGVTHPT;'
        'Trusted_Connection=yes;'
    )
    cursor = conn.cursor()
except Exception as e:
    messagebox.showerror("Lỗi kết nối", f"Không thể kết nối cơ sở dữ liệu:\n{e}")

#===== HÀM CANH GIỮA CỬA SỔ =====
def center_window(win, width, height):
    screen_width = win.winfo_screenwidth()
    screen_height = win.winfo_screenheight()
    x = (screen_width // 2) - (width // 2)
    y = (screen_height // 2) - (height // 2)
    win.geometry(f'{width}x{height}+{x}+{y}')

#===== CỬA SỔ ĐĂNG NHẬP =====
root = Tk()
root.title("Đăng nhập hệ thống")
center_window(root, 925, 500)
root.configure(bg="#fff")
root.resizable(False, False)

#ĐĂNG NHẬP
tai_khoan = {"admin": "1234"}  
def dang_nhap():
    ten_dang_nhap = nhap_ten.get()
    mat_khau = nhap_mk.get()
    if ten_dang_nhap in tai_khoan and tai_khoan[ten_dang_nhap] == mat_khau:
        root.withdraw()
        open_main_window()
    else:
        messagebox.showerror("Lỗi đăng nhập", "Tên đăng nhập hoặc mật khẩu không đúng!")

#GIAO DIỆN ĐĂNG NHẬP
img = PhotoImage(file='login.png')
Label(root, image=img, bg='white').place(x=50, y=50)
khung = Frame(root, width=350, height=350, bg="white")
khung.place(x=480, y=70)
tieu_de = Label(khung, text='ĐĂNG NHẬP', fg='#57a1f8', bg='white',
                font=('Times New Roman', 23, 'bold'))
tieu_de.place(x=85, y=5)

def on_enter(e):
    if nhap_ten.get() == 'Tên đăng nhập':
        nhap_ten.delete(0, 'end')
def on_leave(e):
    if nhap_ten.get() == '':
        nhap_ten.insert(0, 'Tên đăng nhập')
nhap_ten = Entry(khung, width=25, fg='black', border=0, bg='white',
                 font=('Times New Roman', 11))
nhap_ten.place(x=30, y=80)
nhap_ten.insert(0, 'Tên đăng nhập')
nhap_ten.bind('<FocusIn>', on_enter)
nhap_ten.bind('<FocusOut>', on_leave)
Frame(khung, width=295, height=2, bg='black').place(x=25, y=107)

def on_enter(e):
    if nhap_mk.get() == 'Mật khẩu':
        nhap_mk.delete(0, 'end')
        nhap_mk.config(show='*')
def on_leave(e):
    if nhap_mk.get() == '':
        nhap_mk.insert(0, 'Mật khẩu')
        nhap_mk.config(show='')
nhap_mk = Entry(khung, width=25, fg='black', border=0, bg='white',
                font=('Times New Roman', 11))
nhap_mk.place(x=30, y=150)
nhap_mk.insert(0, 'Mật khẩu')
nhap_mk.bind('<FocusIn>', on_enter)
nhap_mk.bind('<FocusOut>', on_leave)
Frame(khung, width=295, height=2, bg='black').place(x=25, y=177)

Button(khung, width=39, pady=7, text='Đăng nhập', bg='#57a1f8',
       fg='white', border=0, command=dang_nhap,
    font=('Times New Roman', 10, 'bold')).place(x=35, y=204)
Label(khung, text="Chưa có tài khoản?", fg='black', bg='white',
      font=('Times New Roman', 9)).place(x=90, y=270)

def mo_dang_ky():
    cua_so_dang_ky()
nut_dangky = Button(khung, width=6, text='Đăng ký', border=0, bg='white',
                    cursor='hand2', fg='#57a1f8', font=('Times New Roman', 9, 'bold'),
                    command=mo_dang_ky)
nut_dangky.place(x=220, y=270)

#===== CỬA SỔ ĐĂNG KÝ =====
def cua_so_dang_ky():
    dangky = Toplevel(root)
    dangky.title("Đăng ký tài khoản")
    center_window(dangky, 925, 500)
    dangky.configure(bg="white")
    dangky.resizable(False, False)
    #ĐĂNG KÝ
    def dang_ky_tk():
        ten = nhap_ten_dk.get()
        mk = nhap_mk_dk.get()
        nhaplai = nhap_xn_mk.get()

        if mk != nhaplai:
            messagebox.showerror("Lỗi", "Mật khẩu xác nhận không khớp!")
        elif ten == "" or mk == "":
            messagebox.showerror("Lỗi", "Vui lòng nhập đầy đủ thông tin!")
        elif ten in tai_khoan:
            messagebox.showerror("Lỗi", "Tên đăng nhập đã tồn tại!")
        else:
            tai_khoan[ten] = mk
            messagebox.showinfo("Thành công", "Đăng ký tài khoản thành công!")
            dangky.destroy()

   #GIAO DIỆN ĐĂNG KÝ
    khung_dk = Frame(dangky, width=350, height=390, bg="white")
    khung_dk.place(x=480, y=50)
    Label(khung_dk, text='ĐĂNG KÝ', fg='#57a1f8', bg='white',
          font=('Times New Roman', 23, 'bold')).place(x=90, y=5)
  
    def on_enter(e):
        if nhap_ten_dk.get() == 'Tên đăng nhập':
            nhap_ten_dk.delete(0, 'end')
    def on_leave(e):
        if nhap_ten_dk.get() == '':
            nhap_ten_dk.insert(0, 'Tên đăng nhập')
    nhap_ten_dk = Entry(khung_dk, width=25, fg='black', border=0, bg='white',
                        font=('Times New Roman', 11))
    nhap_ten_dk.place(x=30, y=80)
    nhap_ten_dk.insert(0, 'Tên đăng nhập')
    nhap_ten_dk.bind('<FocusIn>', on_enter)
    nhap_ten_dk.bind('<FocusOut>', on_leave)
    Frame(khung_dk, width=295, height=2, bg='black').place(x=25, y=107)

    def on_enter(e):
        if nhap_mk_dk.get() == 'Mật khẩu':
            nhap_mk_dk.delete(0, 'end')
            nhap_mk_dk.config(show='*')
    def on_leave(e):
        if nhap_mk_dk.get() == '':
            nhap_mk_dk.insert(0, 'Mật khẩu')
            nhap_mk_dk.config(show='')
    nhap_mk_dk = Entry(khung_dk, width=25, fg='black', border=0, bg='white',
                       font=('Times New Roman', 11))
    nhap_mk_dk.place(x=30, y=150)
    nhap_mk_dk.insert(0, 'Mật khẩu')
    nhap_mk_dk.bind('<FocusIn>', on_enter)
    nhap_mk_dk.bind('<FocusOut>', on_leave)
    Frame(khung_dk, width=295, height=2, bg='black').place(x=25, y=177)

    def on_enter(e):
        if nhap_xn_mk.get() == 'Nhập lại mật khẩu':
            nhap_xn_mk.delete(0, 'end')
            nhap_xn_mk.config(show='*')
    def on_leave(e):
        if nhap_xn_mk.get() == '':
            nhap_xn_mk.insert(0, 'Nhập lại mật khẩu')
            nhap_xn_mk.config(show='')
    nhap_xn_mk = Entry(khung_dk, width=25, fg='black', border=0, bg='white',
                       font=('Times New Roman', 11))
    nhap_xn_mk.place(x=30, y=220)
    nhap_xn_mk.insert(0, 'Nhập lại mật khẩu')
    nhap_xn_mk.bind('<FocusIn>', on_enter)
    nhap_xn_mk.bind('<FocusOut>', on_leave)
    Frame(khung_dk, width=295, height=2, bg='black').place(x=25, y=247)

    Button(khung_dk, width=39, pady=7, text='Đăng ký', bg='#57a1f8',
           fg='white', border=0, font=('Times New Roman', 10, 'bold'),
           command=dang_ky_tk).place(x=35, y=280)

    Label(khung_dk, text='Đã có tài khoản?', fg='black', bg='white',
          font=('Times New Roman', 9)).place(x=90, y=340)

    Button(khung_dk, width=8, text='Đăng nhập', border=0, bg='white',
           cursor='hand2', fg='#57a1f8', font=('Times New Roman', 9, 'bold'),
           command=dangky.destroy).place(x=200, y=340)

#===== CỬA SỔ MAIN =====
def open_main_window():
    main = Toplevel(root)
    main.title("Quản lý giáo viên THPT")
    center_window(main, 1200, 700)
    main.configure(bg="white")
    main.resizable(True, True)
    
    #Phần menu
    def dang_xuat():
        if messagebox.askyesno("Đăng xuất", "Bạn có chắc chắn muốn đăng xuất không?"):
            main.destroy()
            root.deiconify()

    def thoat_chuong_trinh():
        if messagebox.askyesno("Thoát", 
            "Bạn có chắc chắn muốn thoát chương trình không?"):
            main.quit()
            main.destroy()

    def mo_link_huong_dan():
        link = "https://sites.google.com/view/hdsdpmqlgvthpt/trang-ch%E1%BB%A7"
        try:
            webbrowser.open_new_tab(link)
        except Exception as e:
            # Sử dụng messagebox.showerror nếu bạn có import messagebox từ tkinter
            messagebox.showerror("Lỗi",
             f"Không thể mở đường link hướng dẫn.\nLỗi: {e}")

    def lien_he_hotline():     
        messagebox.showinfo(
            "Liên hệ hỗ trợ", 
            "Liên hệ hotline: 1900 1753\n(Hoạt động từ 8h00 - 17h00, T2-T6)"
        )
    # ===== MENU CHÍNH =====
    menubar = Menu(main, font=("Times New Roman", 13))

    # ----- Menu Hệ thống -----
    menu_hethong = Menu(menubar, tearoff=0, font=("Times New Roman", 12))
    menu_hethong.add_command(label="Đăng xuất", command=dang_xuat)
    menu_hethong.add_separator()
    menu_hethong.add_command(label="Thoát", command=thoat_chuong_trinh)
    menubar.add_cascade(label="Hệ thống", menu=menu_hethong)

    menu_trogiup = Menu(menubar, tearoff=0, font=("Times New Roman", 12))
    menu_trogiup.add_command(label="Hướng dẫn",command=mo_link_huong_dan)
    menu_trogiup.add_command(label="Liên hệ hotline",command=lien_he_hotline)
    menubar.add_cascade(label="Trợ giúp", menu=menu_trogiup)

    #Phần thanh chức năng
    main.config(menu=menubar)
    thanh_chucnang = Frame(main, bg="#57a1f8", height=65)
    thanh_chucnang.pack(fill=X)
    style_button = {
        "bg": "#57a1f8",
        "fg": "white",
        "font": ("Times New Roman", 14, "bold"),
        "border": 0,
        "activebackground": "#468be8",
        "activeforeground": "white",
        "cursor": "hand2",
        "pady": 10,
        "width": 15
    }
    def tao_vach_ngan(parent):
        Frame(parent, bg="white", width=2, height=40).pack(side=LEFT, pady=10)

    Button(thanh_chucnang, text="Giáo viên",
       command=lambda: hien_thi_giao_vien(),
       **style_button).pack(side=LEFT)
    tao_vach_ngan(thanh_chucnang)

    Button(thanh_chucnang, text="Lương",
       command=lambda: hien_thi_luong(),
       **style_button).pack(side=LEFT)
    tao_vach_ngan(thanh_chucnang)

    Button(thanh_chucnang, text="Hợp đồng",
       command=lambda: hien_thi_hopdong(),
       **style_button).pack(side=LEFT)
    tao_vach_ngan(thanh_chucnang)

    Button(thanh_chucnang, text="Nghỉ phép",
       command=lambda: hien_thi_nghi_phep(),
       **style_button).pack(side=LEFT)
    tao_vach_ngan(thanh_chucnang)
    

    #=====QUẢN LÝ GIÁO VIÊN =====
    noi_dung_frame = Frame(main, bg="white")
    noi_dung_frame.pack(fill=BOTH, expand=True)

    #Tải dữ liệu giáo viên
    def load_data_giaovien():
        global tree
        for row in tree.get_children():
            tree.delete(row)
        try:
            cursor.execute("""
                SELECT gv.MaGV, gv.HoTenGV, gv.GioiTinh,
                 gv.NgaySinh, gv.SoDT, gv.Email, gv.CMND,
                 gv.DiaChi,cv.TenCV, tcm.TenToCM
                FROM GiaoVien gv
                LEFT JOIN ChucVu cv ON gv.MaCV = cv.MaCV
                LEFT JOIN ToChuyenMon tcm ON gv.MaToCM = tcm.MaToCM
            """)
            rows = cursor.fetchall()
            for row in rows:
                tree.insert("", "end", values=[str(col) for col in row])

        except Exception as e:
            messagebox.showerror("Lỗi truy vấn",
             f"Không thể tải dữ liệu giáo viên:\n{e}")

    #HIỂN THỊ GIÁO VIÊN
    def hien_thi_giao_vien():
        global tree
        for widget in noi_dung_frame.winfo_children():
            widget.destroy()

        Label(noi_dung_frame, text="QUẢN LÝ GIÁO VIÊN",
            font=("Times New Roman", 24, "bold"),
            fg="#0a46a3", bg="white").pack(pady=10)

        khung_nhap = LabelFrame(noi_dung_frame, text="Nhập thông tin",
                                font=("Times New Roman", 14, "bold"),
                                bg="#cce4f7", fg="#0a46a3", padx=15, pady=10)
        khung_nhap.pack(fill=X, padx=20, pady=10)

        search_frame = Frame(khung_nhap, bg="#cce4f7")
        search_frame.grid(row=0, column=0, columnspan=5, sticky=W, pady=5)

        global entry_timkiem
        Label(search_frame, text="Tìm kiếm:", bg="#cce4f7",
         font=("Times New Roman", 13)).pack(side=LEFT, padx=5)
        entry_timkiem = Entry(search_frame, width=30, font=("Times New Roman", 13))
        entry_timkiem.pack(side=LEFT, padx=5)
        placeholder = "Nhập mã giáo viên..."
        entry_timkiem.insert(0, placeholder)
        entry_timkiem.config(fg="grey")
        def on_focus_in(event):
            if entry_timkiem.get() == placeholder:
                entry_timkiem.delete(0, END)
                entry_timkiem.config(fg="black")
        def on_focus_out(event):
            if entry_timkiem.get().strip() == "":
                entry_timkiem.insert(0, placeholder)
                entry_timkiem.config(fg="grey")
        entry_timkiem.bind("<FocusIn>", on_focus_in)
        entry_timkiem.bind("<FocusOut>", on_focus_out)
        
        global entry_magv, entry_hoten, entry_ngaysinh, entry_sdt
        global entry_cccd, entry_email, entry_diachi, entry_chucvu, entry_bomon, gioitinh_var
        Label(khung_nhap, text="Mã giáo viên:", bg="#cce4f7",
               font=("Times New Roman", 13)).grid(row=1, column=0, sticky=W, padx=5, pady=5)
        entry_magv = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_magv.grid(row=1, column=1, padx=5, pady=5)

        Label(khung_nhap, text="Họ tên:", bg="#cce4f7", 
              font=("Times New Roman", 13)).grid(row=1, column=2, sticky=W, padx=5, pady=5)
        entry_hoten = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_hoten.grid(row=1, column=3, padx=5, pady=5)

        Label(khung_nhap, text="Giới tính:", bg="#cce4f7", 
              font=("Times New Roman", 13)).grid(row=2, column=0, sticky=W, padx=5, pady=5)
        gioitinh_var = StringVar(value="Nam")  
        Frame_gioitinh = Frame(khung_nhap, bg="#cce4f7")
        Frame_gioitinh.grid(row=2, column=1, sticky=W, padx=5, pady=5)
        Radiobutton(Frame_gioitinh, text="Nam", variable=gioitinh_var, 
               value="Nam", bg="#cce4f7", font=("Times New Roman", 13)).pack(side=LEFT, padx=10)
        Radiobutton(Frame_gioitinh, text="Nữ", variable=gioitinh_var, 
               value="Nữ", bg="#cce4f7", font=("Times New Roman", 13)).pack(side=LEFT)

        Label(khung_nhap, text="Ngày sinh:", bg="#cce4f7", 
              font=("Times New Roman", 13)).grid(row=2, column=2, sticky=W, padx=5, pady=5)
        entry_ngaysinh = DateEntry(khung_nhap, width=23, font=("Times New Roman", 13), 
              background='darkblue', foreground='white', borderwidth=2, date_pattern='yyyy-mm-dd')
        entry_ngaysinh.grid(row=2, column=3, padx=5, pady=5)

        Label(khung_nhap, text="Số điện thoại:", bg="#cce4f7", 
              font=("Times New Roman", 13)).grid(row=3, column=0, sticky=W, padx=5, pady=5)
        entry_sdt = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_sdt.grid(row=3, column=1, padx=5, pady=5)

        Label(khung_nhap, text="CCCD:", bg="#cce4f7", 
              font=("Times New Roman", 13)).grid(row=3, column=2, sticky=W, padx=5, pady=5)
        entry_cccd = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_cccd.grid(row=3, column=3, padx=5, pady=5)

        Label(khung_nhap, text="Email:", bg="#cce4f7", 
              font=("Times New Roman", 13)).grid(row=4, column=0, sticky=W, padx=5, pady=5)
        entry_email = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_email.grid(row=4, column=1, padx=5, pady=5)

        Label(khung_nhap, text="Địa chỉ:", bg="#cce4f7", 
              font=("Times New Roman", 13)).grid(row=4, column=2, sticky=W, padx=5, pady=5)
        entry_diachi = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_diachi.grid(row=4, column=3, padx=5, pady=5)

        Label(khung_nhap, text="Bộ môn:", bg="#cce4f7", 
              font=("Times New Roman", 13)).grid(row=5, column=0, sticky=W, padx=5, pady=5)
        cursor.execute("SELECT TenToCM FROM ToChuyenMon")
        bomon_list = [row[0] for row in cursor.fetchall()]
        entry_bomon = ttk.Combobox(khung_nhap, values=bomon_list, 
                                   font=("Times New Roman", 13), width=23)
        entry_bomon.grid(row=5, column=1, padx=5, pady=5)
        entry_bomon.config(state="readonly")

        Label(khung_nhap, text="Chức vụ:", bg="#cce4f7", 
              font=("Times New Roman", 13)).grid(row=5, column=2, sticky=W, padx=5, pady=5)
        cursor.execute("SELECT TenCV FROM ChucVu")
        cv_list = [row[0] for row in cursor.fetchall()]
        entry_chucvu = ttk.Combobox(khung_nhap, values=cv_list, 
                                    font=("Times New Roman", 13), width=23)
        entry_chucvu.grid(row=5, column=3, padx=5, pady=5)
        entry_chucvu.config(state="readonly")

        nut_frame = Frame(khung_nhap, bg="#cce4f7")
        nut_frame.grid(row=1, column=4, rowspan=5, padx=20, sticky=N)
        #Các CRUD giáo viên
        #Hiển thị entry
        def hien_thi_entry(event):
            global tree
            selected = tree.focus()  
            if not selected:
                return
            values = tree.item(selected, 'values')
            entry_magv.delete(0, END)
            entry_magv.insert(0, values[0])  
            entry_hoten.delete(0, END)
            entry_hoten.insert(0, values[1])       
            gioitinh_var.set(values[2])
            try:
                dt = datetime.strptime(values[3], "%Y-%m-%d")
                entry_ngaysinh.set_date(dt)
            except:
                pass
            entry_sdt.delete(0, END)
            entry_sdt.insert(0, values[4])
            entry_email.delete(0, END)
            entry_email.insert(0, values[5])
            entry_cccd.delete(0, END)
            entry_cccd.insert(0, values[6])
            entry_diachi.delete(0, END)
            entry_diachi.insert(0, values[7])
            entry_chucvu.set(values[8])
            entry_bomon.set(values[9])

        #Thêm
        def them_giao_vien():
            magv = entry_magv.get().strip()
            hoten = entry_hoten.get().strip()
            gioitinh = gioitinh_var.get()
            ngaysinh_dt = entry_ngaysinh.get_date().strftime("%Y-%m-%d")
            sodt = entry_sdt.get().strip()
            cccd = entry_cccd.get().strip()
            email = entry_email.get().strip()
            diachi = entry_diachi.get().strip()
            tencv = entry_chucvu.get().strip()
            tenbomon = entry_bomon.get().strip()
        
            if not magv or not hoten:
                messagebox.showwarning("Thiếu dữ liệu", 
                            "Vui lòng nhập đầy đủ thông tin")
                return
            try:
                cursor.execute("SELECT COUNT(*) FROM GiaoVien WHERE MaGV=?", (magv,))
                if cursor.fetchone()[0] > 0:
                    messagebox.showwarning("Trùng mã", 
                                "Mã giáo viên đã tồn tại, vui lòng nhập mã khác!")
                    return
                cursor.execute("SELECT MaCV FROM ChucVu WHERE TenCV=?", (tencv,))
                result_cv = cursor.fetchone()
                if result_cv:
                    macv = result_cv[0]
                else:
                    messagebox.showwarning("Chức vụ chưa có", 
                                f"Chức vụ '{tencv}' không tồn tại trong hệ thống!")
                    return
                cursor.execute("SELECT MaToCM FROM ToChuyenMon WHERE TenToCM=?", (tenbomon,))
                result_tm = cursor.fetchone()
                if result_tm:
                    matocm = result_tm[0]
                else:
                    messagebox.showwarning("Bộ môn chưa có",
                             f"Bộ môn '{tenbomon}' không tồn tại trong hệ thống!")
                    return
                cursor.execute("""
                        INSERT INTO GiaoVien (MaGV, HoTenGV, GioiTinh,
                             NgaySinh, SoDT, Email, CMND, DiaChi, MaCV, MaToCM)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, (magv, hoten, gioitinh, ngaysinh_dt, sodt, email, cccd, diachi, macv, matocm))
                conn.commit()
                messagebox.showinfo("Thành công", "Đã thêm giáo viên mới!")
                load_data_giaovien()
            except Exception as e:
                messagebox.showerror("Lỗi", f"Không thể thêm giáo viên:\n{e}")
        Button(nut_frame, text="Thêm", bg="#0D47A1", fg="white",
            font=("Times New Roman", 13, "bold"), width=10,command=lambda: them_giao_vien()).pack(pady=5)

        #Sửa
        def sua_giao_vien():
            magv = entry_magv.get().strip()
            hoten = entry_hoten.get().strip()
            gioitinh = gioitinh_var.get()
            ngaysinh_dt = entry_ngaysinh.get_date().strftime("%Y-%m-%d")
            sodt = entry_sdt.get().strip()
            cccd = entry_cccd.get().strip()
            email = entry_email.get().strip()
            diachi = entry_diachi.get().strip()
            tencv = entry_chucvu.get().strip()
            tenbomon = entry_bomon.get().strip()
            
            if not magv or not hoten:
                messagebox.showwarning("Thiếu dữ liệu", "Vui lòng nhập đầy đủ thông tin")
                return
            try:
                cursor.execute("SELECT COUNT(*) FROM GiaoVien WHERE MaGV=?", (magv,))
                if cursor.fetchone()[0] == 0:
                    messagebox.showwarning("Không tồn tại", "Mã giáo viên này không tồn tại!")
                    return
                cursor.execute("SELECT MaCV FROM ChucVu WHERE TenCV=?", (tencv,))
                result_cv = cursor.fetchone()
                if result_cv:
                    macv = result_cv[0]
                else:
                    messagebox.showwarning("Chức vụ chưa có",
                                            f"Chức vụ '{tencv}' không tồn tại trong hệ thống!")
                    return
                cursor.execute("SELECT MaToCM FROM ToChuyenMon WHERE TenToCM=?", (tenbomon,))
                result_tm = cursor.fetchone()
                if result_tm:
                    matocm = result_tm[0]
                else:
                    messagebox.showwarning("Bộ môn chưa có",
                                            f"Bộ môn '{tenbomon}' không tồn tại trong hệ thống!")
                    return
                cursor.execute("""
                    UPDATE GiaoVien
                    SET HoTenGV=?, GioiTinh=?, NgaySinh=?, SoDT=?, Email=?, CMND=?, DiaChi=?, MaCV=?, MaToCM=?
                    WHERE MaGV=?
                """, (hoten, gioitinh, ngaysinh_dt, sodt, email, cccd, diachi, macv, matocm, magv))
                conn.commit()
                messagebox.showinfo("Thành công", "Đã sửa thông tin giáo viên!")
                load_data_giaovien()
            except Exception as e:
                messagebox.showerror("Lỗi", f"Không thể sửa giáo viên:\n{e}")
        Button(nut_frame, text="Sửa", bg="#0D47A1", fg="white",
                font=("Times New Roman", 13, "bold"), width=10,command=lambda: sua_giao_vien()).pack(pady=5)

        #Làm mới các entry
        def lam_moi_form():
            entry_magv.delete(0, END)
            entry_hoten.delete(0, END)
            gioitinh_var.set("")  
            try:
                entry_ngaysinh.set_date(datetime.now())  
            except:
                pass
            entry_sdt.delete(0, END)
            entry_cccd.delete(0, END)
            entry_email.delete(0, END)
            entry_diachi.delete(0, END)
            entry_chucvu.set("")   
            entry_bomon.set("")    
            entry_magv.focus()
        Button(nut_frame, text="Làm mới", bg="#0D47A1", fg="white",
               font=("Times New Roman", 13, "bold"), width=10,command=lam_moi_form).pack(pady=5)

        #Xóa
        def xoa_giao_vien():
            selected = tree.focus()
            if not selected:
                messagebox.showwarning("Chưa chọn", "Vui lòng chọn giáo viên cần xóa trong danh sách!")
                return
            values = tree.item(selected, 'values')
            magv = values[0]
            ten = values[1]
            confirm = messagebox.askyesno("Xác nhận xóa", f"Bạn có chắc chắn muốn xóa giáo viên này không?")
            if not confirm:
                return
            try:
                cursor.execute("DELETE FROM GiaoVien WHERE MaGV = ?", (magv,))
                conn.commit()
                messagebox.showinfo("Thành công", f"Đã xóa giáo viên thành công!")
                load_data_giaovien() 
            except Exception as e:
                messagebox.showerror("Lỗi", f"Không thể xóa giáo viên:\n{e}")
        Button(nut_frame, text="Xóa", bg="#0D47A1", fg="white",
                font=("Times New Roman", 13, "bold"), width=10,command=lambda: xoa_giao_vien()).pack(pady=5)

        #Lưu
        def luu_giao_vien():
            if not tree.get_children():
                messagebox.showwarning("Chưa có dữ liệu", "Danh sách giáo viên trống!")
                return
            wb = Workbook()
            ws = wb.active
            ws.title = "Danh sách giáo viên"
            columns = ("Mã GV", "Họ tên", "Giới tính", "Ngày sinh", "SĐT", "Email", "CCCD", "Địa chỉ", "Chức vụ", "Bộ môn")
            ws.append(columns)
            for row_id in tree.get_children():
                values = tree.item(row_id)['values']
                ws.append(values)
            file_path = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                                    filetypes=[("Excel files","*.xlsx")])
            if file_path:
                wb.save(file_path)
                messagebox.showinfo("Thành công", f"Đã lưu file Excel: {file_path}")
        Button(nut_frame, text="Lưu", bg="#0D47A1", fg="white", 
               font=("Times New Roman", 13, "bold"), width=10,command=lambda: luu_giao_vien()).pack(pady=5)

        #Tìm kiếm
        def tim_gv():
            global tree, entry_timkiem
            magv_tim = entry_timkiem.get().strip()
            if magv_tim == "" or magv_tim == "Nhập mã giáo viên...":
                load_data_giaovien()
                return
            for row in tree.get_children():
                tree.delete(row)
            try:
                cursor.execute("""
                    SELECT gv.MaGV, gv.HoTenGV, gv.GioiTinh, 
                            gv.NgaySinh, gv.SoDT, gv.Email, gv.CMND, gv.DiaChi, cv.TenCV, tcm.TenToCM
                    FROM GiaoVien gv
                    LEFT JOIN ChucVu cv ON gv.MaCV = cv.MaCV
                    LEFT JOIN ToChuyenMon tcm ON gv.MaToCM = tcm.MaToCM
                    WHERE gv.MaGV = ?
                """, (magv_tim,))
                row = cursor.fetchone()
                if row:
                    tree.insert("", "end", values=[str(col) for col in row])
                else:
                    messagebox.showinfo("Không tìm thấy", f"Không tìm thấy giáo viên có Mã GV '{magv_tim}'")
                    load_data_giaovien()    
            except Exception as e:
                messagebox.showerror("Lỗi", f"Không thể tìm giáo viên:\n{e}")
        Button(search_frame, text="Tìm", bg="#e80101", fg="white", 
               font=("Times New Roman", 13, "bold"), command=tim_gv).pack(side=LEFT, padx=5)

        #Hiển thị hồ sơ
        def hien_thi_ho_so():
            selected = tree.focus()  # Lấy dòng được chọn
            if not selected:
                messagebox.showwarning("Chưa chọn", "Vui lòng chọn giáo viên cần hiển thị hồ sơ!")
                return
            values = tree.item(selected, 'values')
            ho_so_win = Toplevel()
            ho_so_win.title(f"Hồ sơ giáo viên: {values[1]}")
            ho_so_win.geometry("500x600")
            ho_so_win.configure(bg="white")
            Label(ho_so_win, text="HỒ SƠ GIÁO VIÊN", font=("Times New Roman", 20, "bold"), fg="#0a46a3", bg="white").pack(pady=10)
            try:
                anh_path = values[10] if len(values) > 10 else None
                if anh_path and os.path.exists(anh_path):
                    from PIL import Image, ImageTk
                    img = Image.open(anh_path)
                    img = img.resize((150, 150))
                    photo = ImageTk.PhotoImage(img)
                    lbl_img = Label(ho_so_win, image=photo, bg="white")
                    lbl_img.image = photo  
                    lbl_img.pack(pady=10)
            except:
                pass
            info_frame = Frame(ho_so_win, bg="white")
            info_frame.pack(pady=10, padx=20, fill=BOTH, expand=True)

            labels = ["Mã GV:", "Họ tên:", "Giới tính:", "Ngày sinh:", "Số điện thoại:", 
                    "Email:", "CCCD:", "Địa chỉ:", "Chức vụ:", "Bộ môn:"]
            for i, label in enumerate(labels):
                lbl = Label(info_frame, text=label, font=("Times New Roman", 13, "bold"), bg="white", anchor=W)
                lbl.grid(row=i, column=0, sticky=W, pady=5)
                val = Label(info_frame, text=values[i], font=("Times New Roman", 13), bg="white", anchor=W)
                val.grid(row=i, column=1, sticky=W, pady=5)
            Button(ho_so_win, text="Đóng", font=("Times New Roman", 13, "bold"), bg="#0D47A1", fg="white",
                command=ho_so_win.destroy).pack(pady=20)
        Button(search_frame, text="Hiển thị hồ sơ", bg="#393ced", fg="white",
                font=("Times New Roman", 13, "bold"),command=hien_thi_ho_so).pack(side=LEFT, padx=5)

        table_frame = Frame(noi_dung_frame, bg="white")
        table_frame.pack(fill=BOTH, expand=True, padx=20, pady=10)

        columns = ("magv", "hoten", "gioitinh", "ngaysinh", "sodienthoai", "email", "cccd", "diachi", "chucvu", "bomon")
        global tree
        tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=25)

        tree.heading("magv", text="Mã GV")
        tree.heading("hoten", text="Họ tên")
        tree.heading("gioitinh", text="Giới tính")
        tree.heading("ngaysinh", text="Ngày sinh")
        tree.heading("sodienthoai", text="SĐT")
        tree.heading("email", text="Email")
        tree.heading("cccd", text="CCCD")
        tree.heading("diachi", text="Địa chỉ")
        tree.heading("chucvu", text="Chức vụ")
        tree.heading("bomon", text="Bộ môn")

        for col in columns:
            tree.column(col, width=100, anchor=W)

        tree.pack(fill=BOTH, expand=True)
        tree.bind("<ButtonRelease-1>", hien_thi_entry)
        # --- Tải dữ liệu từ SQL ---
        load_data_giaovien()

    #=====QUẢN LÝ LƯƠNG=====
    noi_dung_frame = Frame(main, bg="white")
    noi_dung_frame.pack(fill=BOTH, expand=True)
    #Tải dữ liệu lương
    def load_data_luong():
        global tree
        for row in tree.get_children():
            tree.delete(row)
        try:
            cursor.execute("""
                SELECT l.MaGV, l.NgayBD, l.LuongCB, l.HeSoLuong, l.PhuCap, l.Thuong, l.GhiChu
                FROM QuaTrinhLuong l
            """)
            rows = cursor.fetchall()
            for row in rows:
                tree.insert("", "end", values=[str(col) for col in row])

        except Exception as e:
            messagebox.showerror("Lỗi truy vấn", f"Không thể tải dữ liệu giáo viên:\n{e}")
    #Hiển thị lương
    def hien_thi_luong():
        global tree
        for widget in noi_dung_frame.winfo_children():
            widget.destroy()

        Label(noi_dung_frame, text="QUẢN LÝ LƯƠNG GIÁO VIÊN",
            font=("Times New Roman", 24, "bold"),
            fg="#0a46a3", bg="white").pack(pady=10)

        khung_nhap = LabelFrame(noi_dung_frame, text="Nhập thông tin",
                                font=("Times New Roman", 14, "bold"),
                                bg="#cce4f7", fg="#0a46a3", padx=15, pady=10)
        khung_nhap.pack(fill=X, padx=20, pady=10)

        search_frame = Frame(khung_nhap, bg="#cce4f7")
        search_frame.grid(row=0, column=0, columnspan=5, sticky=W, pady=5)

        global entry_timkiem
        Label(search_frame, text="Tìm kiếm:", bg="#cce4f7",
               font=("Times New Roman", 13)).pack(side=LEFT, padx=5)
        entry_timkiem = Entry(search_frame, width=30, font=("Times New Roman", 13))
        entry_timkiem.pack(side=LEFT, padx=5)
        placeholder = "Nhập mã giáo viên..."
        entry_timkiem.insert(0, placeholder)
        entry_timkiem.config(fg="grey")
        def on_focus_in(event):
            if entry_timkiem.get() == placeholder:
                entry_timkiem.delete(0, END)
                entry_timkiem.config(fg="black")
        def on_focus_out(event):
            if entry_timkiem.get().strip() == "":
                entry_timkiem.insert(0, placeholder)
                entry_timkiem.config(fg="grey")
        entry_timkiem.bind("<FocusIn>", on_focus_in)
        entry_timkiem.bind("<FocusOut>", on_focus_out)
        
        global entry_magv, entry_ngaybd, entry_luongcb, entry_hsluong, entry_phucap, entry_thuong, entry_ghichu
        Label(khung_nhap, text="Mã giáo viên:", bg="#cce4f7", 
              font=("Times New Roman", 13)).grid(row=1, column=0, sticky=W, padx=5, pady=5)
        entry_magv = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_magv.grid(row=1, column=1, padx=5, pady=5)

        Label(khung_nhap, text="Ngày bắt đầu:", bg="#cce4f7", 
              font=("Times New Roman", 13)).grid(row=1, column=2, sticky=W, padx=5, pady=5)
        entry_ngaybd = DateEntry(khung_nhap, width=23,
            font=("Times New Roman", 13), background='darkblue', foreground='white', borderwidth=2, date_pattern='yyyy-mm-dd')
        entry_ngaybd.grid(row=1, column=3, padx=5, pady=5)

        Label(khung_nhap, text="Lương cơ bản:", bg="#cce4f7",
               font=("Times New Roman", 13)).grid(row=2, column=0, sticky=W, padx=5, pady=5)
        entry_luongcb = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_luongcb.grid(row=2, column=1, padx=5, pady=5)

        Label(khung_nhap, text="Hệ số lương:", bg="#cce4f7",
               font=("Times New Roman", 13)).grid(row=2, column=2, sticky=W, padx=5, pady=5)
        entry_hsluong = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_hsluong.grid(row=2, column=3, padx=5, pady=5)

        Label(khung_nhap, text="Phụ cấp:", bg="#cce4f7", 
              font=("Times New Roman", 13)).grid(row=3, column=0, sticky=W, padx=5, pady=5)
        entry_phucap = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_phucap.grid(row=3, column=1, padx=5, pady=5)

        Label(khung_nhap, text="Thưởng:", bg="#cce4f7",
               font=("Times New Roman", 13)).grid(row=3, column=2, sticky=W, padx=5, pady=5)
        entry_thuong = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_thuong.grid(row=3, column=3, padx=5, pady=5)

        Label(khung_nhap, text="Ghi chú:", bg="#cce4f7",
               font=("Times New Roman", 13)).grid(row=4, column=0, sticky=W, padx=5, pady=5)
        entry_ghichu = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_ghichu.grid(row=4, column=1, padx=5, pady=5)


        nut_frame = Frame(khung_nhap, bg="#cce4f7")
        nut_frame.grid(row=1, column=4, rowspan=5, padx=20, sticky=N)
        #Các CRUD lương
        #Hiển thị entry
        def hien_thi_entry(event):
            global tree
            selected = tree.focus()
            if not selected:
                return
            values = tree.item(selected, 'values')
            entry_magv.delete(0, END)
            entry_magv.insert(0, values[0])
            
            try:
                dt = datetime.strptime(values[1], "%Y-%m-%d")
                entry_ngaybd.set_date(dt)
            except:
                pass
            
            entry_luongcb.delete(0, END)
            entry_luongcb.insert(0, values[2])
            
            entry_hsluong.delete(0, END)
            entry_hsluong.insert(0, values[3])
            
            entry_phucap.delete(0, END)
            entry_phucap.insert(0, values[4])
            
            entry_thuong.delete(0, END)
            entry_thuong.insert(0, values[5])
            
            entry_ghichu.delete(0, END)
            entry_ghichu.insert(0, values[6])

        #Thêm
        def them_luong():
            magv = entry_magv.get().strip()
            ngaybd_dt = entry_ngaybd.get_date().strftime("%Y-%m-%d")
            luongcb = entry_luongcb.get().strip()
            hsluong = entry_hsluong.get().strip()
            phucap = entry_phucap.get().strip()
            thuong = entry_thuong.get().strip()
            ghichu = entry_ghichu.get().strip()

            if not magv or not luongcb or not hsluong:
                messagebox.showwarning("Thiếu dữ liệu", "Vui lòng nhập đầy đủ các trường bắt buộc!")
                return

            try:
                # Kiểm tra xem mã giáo viên có trong bảng GiaoVien không
                cursor.execute("SELECT MaGV FROM GiaoVien WHERE MaGV = ?", (magv,))
                if not cursor.fetchone():
                    messagebox.showerror("Lỗi", f"Mã giáo viên '{magv}' không tồn tại!\nVui lòng nhập lại mã giáo viên khác.")
                    entry_magv.delete(0, END)
                    entry_magv.focus()
                    return

                # Kiểm tra (MaGV, NgayBD) đã tồn tại chưa
                cursor.execute("SELECT * FROM QuaTrinhLuong WHERE MaGV=? AND NgayBD=?", (magv, ngaybd_dt))
                if cursor.fetchone():
                    messagebox.showwarning("Bản ghi tồn tại", f"Lương của giáo viên {magv} ngày {ngaybd_dt} đã tồn tại!")
                    return

                # Thêm mới nếu hợp lệ
                cursor.execute("""
                    INSERT INTO QuaTrinhLuong (MaGV, NgayBD, LuongCB, HeSoLuong, PhuCap, Thuong, GhiChu)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                """, (magv, ngaybd_dt, luongcb, hsluong, phucap if phucap else None,
                    thuong if thuong else None, ghichu if ghichu else None))
                conn.commit()
                messagebox.showinfo("Thành công", "Đã thêm quá trình lương mới!")
                load_data_luong()

            except Exception as e:
                messagebox.showerror("Lỗi", f"Không thể thêm lương:\n{e}")

        Button(nut_frame, text="Thêm", bg="#0D47A1", fg="white",
               font=("Times New Roman", 13, "bold"), width=10,command=lambda: them_luong()).pack(pady=5)

        #Sửa
        def sua_luong():
            magv = entry_magv.get().strip()
            ngaybd_dt = entry_ngaybd.get_date().strftime("%Y-%m-%d")
            luongcb = entry_luongcb.get().strip()
            hsluong = entry_hsluong.get().strip()
            phucap = entry_phucap.get().strip()
            thuong = entry_thuong.get().strip()
            ghichu = entry_ghichu.get().strip()

            if not magv or not ngaybd_dt:
                messagebox.showwarning("Thiếu dữ liệu", "Vui lòng chọn một bản ghi để sửa!")
                return

            try:
                # Kiểm tra mã giáo viên có tồn tại không
                cursor.execute("SELECT MaGV FROM GiaoVien WHERE MaGV = ?", (magv,))
                if not cursor.fetchone():
                    messagebox.showerror("Lỗi",
                         f"Mã giáo viên '{magv}' không tồn tại!\nVui lòng nhập lại mã giáo viên khác.")
                    entry_magv.delete(0, END)
                    entry_magv.focus()
                    return

                # Kiểm tra bản ghi (MaGV, NgayBD) có tồn tại trong bảng QuaTrinhLuong không
                cursor.execute("SELECT * FROM QuaTrinhLuong WHERE MaGV=? AND NgayBD=?", (magv, ngaybd_dt))
                if not cursor.fetchone():
                    messagebox.showwarning("Không tồn tại",
                                 f"Bản ghi lương của giáo viên {magv} ngày {ngaybd_dt} không tồn tại!")
                    return

                # Cập nhật thông tin lương
                cursor.execute("""
                    UPDATE QuaTrinhLuong
                    SET LuongCB=?, HeSoLuong=?, PhuCap=?, Thuong=?, GhiChu=?
                    WHERE MaGV=? AND NgayBD=?
                """, (luongcb, hsluong, phucap if phucap else None,
                    thuong if thuong else None, ghichu if ghichu else None,
                    magv, ngaybd_dt))
                conn.commit()
                messagebox.showinfo("Thành công", "Đã sửa thông tin lương!")
                load_data_luong()

            except Exception as e:
                messagebox.showerror("Lỗi", f"Không thể sửa lương:\n{e}")


        Button(nut_frame, text="Sửa", bg="#0D47A1", fg="white",
                font=("Times New Roman", 13, "bold"), width=10,command=lambda: sua_luong()).pack(pady=5)

        #Làm mới các entry
        def lam_moi_form():
            entry_magv.delete(0, END)
            try:
                entry_ngaybd.set_date(datetime.now())  
            except:
                pass
            entry_luongcb.delete(0, END)
            entry_hsluong.delete(0, END)
            entry_phucap.delete(0, END)
            entry_thuong.delete(0, END)
            entry_ghichu.delete(0, END)
            entry_magv.focus()
        Button(nut_frame, text="Làm mới", bg="#0D47A1", fg="white",
               font=("Times New Roman", 13, "bold"), width=10,command=lam_moi_form).pack(pady=5)

        #Xóa
        def xoa_luong():
            global tree
            selected = tree.focus()
            if not selected:
                messagebox.showwarning("Chưa chọn", "Vui lòng chọn dòng lương cần xóa!")
                return
            values = tree.item(selected, 'values')
            magv, ngaybd = values[0], values[1]
            confirm = messagebox.askyesno("Xác nhận xóa",
                             f"Bạn có chắc chắn muốn xóa lương của {magv} (ngày {ngaybd}) không?")
            if not confirm:
                return

            try:
                cursor.execute("DELETE FROM QuaTrinhLuong WHERE MaGV=? AND NgayBD=?", (magv, ngaybd))
                conn.commit()
                messagebox.showinfo("Thành công", f"Đã xóa lương của {magv} ngày {ngaybd}")
                load_data_luong()
            except Exception as e:
                messagebox.showerror("Lỗi", f"Không thể xóa lương:\n{e}")

        Button(nut_frame, text="Xóa", bg="#0D47A1", fg="white", 
               font=("Times New Roman", 13, "bold"), width=10,command=lambda: xoa_luong()).pack(pady=5)

        #Lưu
        def luu_luong():
            if not tree.get_children():
                messagebox.showwarning("Chưa có dữ liệu", "Danh sách giáo viên trống!")
                return
            wb = Workbook()
            ws = wb.active
            ws.title = "Danh sách lương"
            columns = ("Mã GV", "Ngày bắt đầu", "Lương cơ bản", "Hệ số lương", "Phụ cấp", "Thưởng", "Ghi chú")
            ws.append(columns)
            for row_id in tree.get_children():
                values = tree.item(row_id)['values']
                ws.append(values)
            file_path = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                                    filetypes=[("Excel files","*.xlsx")])
            if file_path:
                wb.save(file_path)
                messagebox.showinfo("Thành công", f"Đã lưu file Excel: {file_path}")
        Button(nut_frame, text="Lưu", bg="#0D47A1", fg="white",
                font=("Times New Roman", 13, "bold"), width=10,command=lambda: luu_luong()).pack(pady=5)

        #Tìm kiếm
        def tim_luong():
            global tree, entry_timkiem
            magv_tim = entry_timkiem.get().strip()
            if magv_tim == "" or magv_tim == "Nhập mã giáo viên...":
                load_data_luong()
                return
            for row in tree.get_children():
                tree.delete(row)
            try:
                cursor.execute("""
                    SELECT l.MaGV, l.NgayBD, l.LuongCB, l.HeSoLuong, l.PhuCap, l.Thuong, l.GhiChu
                    FROM QuaTrinhLuong l
                    WHERE l.MaGV = ?
                """, (magv_tim,))
                rows = cursor.fetchall()
                if rows:
                    for row in rows:
                        tree.insert("", "end", values=[str(col) for col in row])
                else:
                    messagebox.showinfo("Không tìm thấy",
                                         f"Không tìm thấy lương của giáo viên có Mã GV '{magv_tim}'")
                    load_data_luong()
            except Exception as e:
                messagebox.showerror("Lỗi", f"Không thể tìm lương:\n{e}")
        Button(search_frame, text="Tìm", bg="#e80101", fg="white", 
               font=("Times New Roman", 13, "bold"), command=tim_luong).pack(side=LEFT, padx=5)

        table_frame = Frame(noi_dung_frame, bg="white")
        table_frame.pack(fill=BOTH, expand=True, padx=20, pady=10)
    
        columns = ("magv", "ngaybd", "luongcb", "hsluong", "phucap", "thuong", "ghichu")
        global tree
        tree = ttk.Treeview(table_frame,columns=columns,show="headings",height=25)

        tree.heading("magv", text="Mã GV")
        tree.heading("ngaybd", text="Ngày bắt đầu")
        tree.heading("luongcb", text="Lương cơ bản")
        tree.heading("hsluong", text="Hệ số lương")
        tree.heading("phucap", text="Phụ cấp")
        tree.heading("thuong", text="Thưởng")
        tree.heading("ghichu", text="Ghi chú")

        for col in columns:
            tree.column(col, width=100, anchor=W)

        tree.pack(fill=BOTH, expand=True)
        tree.bind("<ButtonRelease-1>", hien_thi_entry)
        # --- Tải dữ liệu từ SQL ---
        load_data_luong()
    
    # ===== QUẢN LÝ HỢP ĐỒNG =====
    def load_data_hopdong():
        global tree_hd
        for row in tree_hd.get_children():
            tree_hd.delete(row)
        try:
            cursor.execute("""
                SELECT h.MaHD, h.NgayBD, h.NgayKT, g.HoTenGV
                FROM HopDong h
                JOIN GiaoVien g ON h.MaGV = g.MaGV
            """)
            rows = cursor.fetchall()
            for row in rows:
                tree_hd.insert("", "end", values=[str(col) for col in row])
        except Exception as e:
            messagebox.showerror("Lỗi truy vấn", f"Không thể tải dữ liệu hợp đồng:\n{e}")

    def hien_thi_hopdong():
        global tree_hd
        for widget in noi_dung_frame.winfo_children():
            widget.destroy()
        Label(noi_dung_frame, text="QUẢN LÝ HỢP ĐỒNG",
            font=("Times New Roman", 24, "bold"),
            fg="#0a46a3", bg="white").pack(pady=10)

        khung_nhap = LabelFrame(noi_dung_frame, text="Nhập thông tin",
                                font=("Times New Roman", 14, "bold"),
                                bg="#cce4f7", fg="#0a46a3", padx=15, pady=10)
        khung_nhap.pack(fill=X, padx=20, pady=10)

        # Entry tìm kiếm
        search_frame = Frame(khung_nhap, bg="#cce4f7")
        search_frame.grid(row=0, column=0, columnspan=4, sticky=W, pady=5)

        global entry_timkiem_hd
        Label(search_frame, text="Tìm kiếm:", bg="#cce4f7",
               font=("Times New Roman", 13)).pack(side=LEFT, padx=5)
        entry_timkiem_hd = Entry(search_frame, width=30, font=("Times New Roman", 13))
        entry_timkiem_hd.pack(side=LEFT, padx=5)
        placeholder = "Nhập mã hợp đồng..."
        entry_timkiem_hd.insert(0, placeholder)
        entry_timkiem_hd.config(fg="grey")
        def on_focus_in(event):
            if entry_timkiem_hd.get() == placeholder:
                entry_timkiem_hd.delete(0, END)
                entry_timkiem_hd.config(fg="black")
        def on_focus_out(event):
            if entry_timkiem_hd.get().strip() == "":
                entry_timkiem_hd.insert(0, placeholder)
                entry_timkiem_hd.config(fg="grey")
        entry_timkiem_hd.bind("<FocusIn>", on_focus_in)
        entry_timkiem_hd.bind("<FocusOut>", on_focus_out)

        # Nút tìm kiếm
        def tim_hopdong():
            mahd_tim = entry_timkiem_hd.get().strip()
            if mahd_tim == "" or mahd_tim == placeholder:
                load_data_hopdong()
                return
            for row in tree_hd.get_children():
                tree_hd.delete(row)
            try:
                cursor.execute("""
                    SELECT h.MaHD, h.NgayBD, h.NgayKT, g.HoTenGV
                    FROM HopDong h
                    JOIN GiaoVien g ON h.MaGV = g.MaGV
                    WHERE h.MaHD=?
                """, (mahd_tim,))
                rows = cursor.fetchall()
                if rows:
                    for row in rows:
                        tree_hd.insert("", "end", values=[str(col) for col in row])
                else:
                    messagebox.showinfo("Không tìm thấy",
                                         f"Không tìm thấy hợp đồng '{mahd_tim}'")
                    load_data_hopdong()
            except Exception as e:
                messagebox.showerror("Lỗi", f"Không thể tìm:\n{e}")
        Button(search_frame, text="Tìm", bg="#e80101",
                fg="white", font=("Times New Roman", 13, "bold"),
            command=tim_hopdong).pack(side=LEFT, padx=5)
        
        # Nút hiển thị hợp đồng nằm ngang với nút tìm kiếm
        def mo_cua_so_hd():
            selected = tree_hd.focus()
            if not selected:
                messagebox.showwarning("Chưa chọn", "Vui lòng chọn hợp đồng để hiển thị!")
                return
            values = tree_hd.item(selected, 'values')
            top = Toplevel()
            top.title(f"Hợp đồng {values[0]}")
            top.geometry("600x500")
            # Nội dung hợp đồng
            text = Text(top, wrap=WORD, font=("Times New Roman", 13))
            text.pack(fill=BOTH, expand=True, padx=10, pady=10)
            
            noi_dung = f"""
                        CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM
                                 Độc lập – Tự do – Hạnh phúc
                                    
                                     HỢP ĐỒNG LAO ĐỘNG
                                            ********
        Hôm nay, ngày {values[1]}, tại trường THPT XYZ, chúng tôi gồm:

        Bên A (Nhà trường): Đại diện…  
        Bên B (Giáo viên): {values[3]}, mã hợp đồng {values[0]}

        Điều 1: Thời hạn hợp đồng
        - Bắt đầu: {values[1]}
        - Kết thúc: {values[2]}

        Điều 2: Nội dung công việc
        - Giảng dạy môn học được phân công
        - Thực hiện các nhiệm vụ khác theo quy định nhà trường

        Điều 3: Quyền lợi và nghĩa vụ
        - Lương, thưởng, phụ cấp theo quy định
        - Thực hiện nghiêm túc các nội quy của nhà trường

        Điều 4: Điều khoản chung
        - Hai bên cam kết thực hiện đúng các điều khoản trên
        - Mọi tranh chấp sẽ được giải quyết theo quy định pháp luật

        Người lập hợp đồng                              Giáo viên
        (Bên A)                                         (Bên B)
            
            """
            text.insert(END, noi_dung)
            text.config(state=DISABLED)
            # Nút lưu hợp đồng
            def luu_hd_word(values, noi_dung):
                try:
                    doc = Document()
                    p = doc.add_paragraph(noi_dung)
                    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
                    if p.runs:
                        p.runs[0].font.size = Pt(12)

                    # --- Lưu file ---
                    file_path = filedialog.asksaveasfilename(
                        defaultextension=".docx",
                        filetypes=[("Word files", "*.docx")],
                        title="Lưu hợp đồng"
                    )
                    if file_path:
                        doc.save(file_path)
                        messagebox.showinfo("Thành công", f"Đã lưu hợp đồng: {file_path}")
                except Exception as e:
                    messagebox.showerror("Lỗi", f"Lưu hợp đồng thất bại!\n{str(e)}")
            Button(top, text="Lưu hợp đồng", bg="#0D47A1", fg="white",
                    font=("Times New Roman", 13, "bold"),
                    command=lambda: luu_hd_word(values, noi_dung)).pack(pady=5)
        Button(search_frame, text="Hiển thị hợp đồng", bg="#0D47A1",
                fg="white", font=("Times New Roman", 13, "bold"),command=mo_cua_so_hd).pack(side=LEFT, padx=5)
            

        # Entry nhập liệu hợp đồng
        global entry_mahd, entry_ngaybd, entry_ngaykt, entry_tengv_hd
        Label(khung_nhap, text="Mã hợp đồng:", bg="#cce4f7", 
              font=("Times New Roman", 13)).grid(row=1, column=0, padx=5, pady=5)
        entry_mahd = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_mahd.grid(row=1, column=1, padx=5, pady=5)

        Label(khung_nhap, text="Ngày bắt đầu:", bg="#cce4f7",
               font=("Times New Roman", 13)).grid(row=1, column=2, padx=5, pady=5)
        entry_ngaybd = DateEntry(khung_nhap, width=23, font=("Times New Roman", 13), date_pattern='yyyy-mm-dd')
        entry_ngaybd.grid(row=1, column=3, padx=5, pady=5)

        Label(khung_nhap, text="Ngày kết thúc:", bg="#cce4f7",
               font=("Times New Roman", 13)).grid(row=2, column=0, padx=5, pady=5)
        entry_ngaykt = DateEntry(khung_nhap, width=23, font=("Times New Roman", 13), date_pattern='yyyy-mm-dd')
        entry_ngaykt.grid(row=2, column=1, padx=5, pady=5)

        Label(khung_nhap, text="Tên giáo viên:", bg="#cce4f7", 
              font=("Times New Roman", 13)).grid(row=2, column=2, padx=5, pady=5)
        entry_tengv_hd = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_tengv_hd.grid(row=2, column=3, padx=5, pady=5)

        # Nút CRUD
        nut_frame = Frame(khung_nhap, bg="#cce4f7")
        nut_frame.grid(row=1, column=4, rowspan=5, padx=20, sticky=N)

        def hien_thi_entry_hd(event):
            selected = tree_hd.focus()
            if not selected:
                return
            values = tree_hd.item(selected, 'values')
            entry_mahd.delete(0, END)
            entry_mahd.insert(0, values[0])
            try:
                entry_ngaybd.set_date(datetime.strptime(values[1], "%Y-%m-%d"))
                entry_ngaykt.set_date(datetime.strptime(values[2], "%Y-%m-%d"))
            except:
                pass
            entry_tengv_hd.delete(0, END)
            entry_tengv_hd.insert(0, values[3])

        # Thêm hợp đồng
        def them_hopdong():
            mahd = entry_mahd.get().strip()
            ngaybd = entry_ngaybd.get_date().strftime("%Y-%m-%d")
            ngaykt = entry_ngaykt.get_date().strftime("%Y-%m-%d")
            tengv = entry_tengv_hd.get().strip()
            if not mahd or not tengv:
                messagebox.showwarning("Thiếu dữ liệu", "Vui lòng nhập đầy đủ các trường bắt buộc!")
                return
            try:
                cursor.execute("SELECT MaGV FROM GiaoVien WHERE HoTenGV=?", (tengv,))
                result = cursor.fetchone()
                if not result:
                    messagebox.showerror("Lỗi", f"Giáo viên '{tengv}' không tồn tại!")
                    return
                magv = result[0]
                cursor.execute("SELECT * FROM HopDong WHERE MaHD=?", (mahd,))
                if cursor.fetchone():
                    messagebox.showwarning("Bản ghi tồn tại", f"Hợp đồng '{mahd}' đã tồn tại!")
                    return
                cursor.execute("INSERT INTO HopDong(MaHD, NgayBD, NgayKT, MaGV) VALUES(?,?,?,?)", (mahd, ngaybd, ngaykt, magv))
                conn.commit()
                messagebox.showinfo("Thành công", "Đã thêm hợp đồng mới!")
                load_data_hopdong()
            except Exception as e:
                messagebox.showerror("Lỗi", f"Không thể thêm:\n{e}")

        Button(nut_frame, text="Thêm", bg="#0D47A1", fg="white", 
               font=("Times New Roman", 13, "bold"), width=10, command=them_hopdong).pack(pady=5)

        # Sửa hợp đồng
        def sua_hopdong():
            mahd = entry_mahd.get().strip()
            ngaybd = entry_ngaybd.get_date().strftime("%Y-%m-%d")
            ngaykt = entry_ngaykt.get_date().strftime("%Y-%m-%d")
            tengv = entry_tengv_hd.get().strip()
            if not mahd:
                messagebox.showwarning("Thiếu dữ liệu", "Vui lòng chọn một hợp đồng để sửa!")
                return
            try:
                cursor.execute("SELECT MaGV FROM GiaoVien WHERE HoTenGV=?", (tengv,))
                result = cursor.fetchone()
                if not result:
                    messagebox.showerror("Lỗi", f"Giáo viên '{tengv}' không tồn tại!")
                    return
                magv = result[0]
                cursor.execute("UPDATE HopDong SET NgayBD=?, NgayKT=?, MaGV=? WHERE MaHD=?", (ngaybd, ngaykt, magv, mahd))
                conn.commit()
                messagebox.showinfo("Thành công", "Đã sửa hợp đồng!")
                load_data_hopdong()
            except Exception as e:
                messagebox.showerror("Lỗi", f"Không thể sửa:\n{e}")

        Button(nut_frame, text="Sửa", bg="#0D47A1", fg="white",
                font=("Times New Roman", 13, "bold"), width=10, command=sua_hopdong).pack(pady=5)

        # Xóa hợp đồng
        def xoa_hopdong():
            selected = tree_hd.focus()
            if not selected:
                messagebox.showwarning("Chưa chọn", "Vui lòng chọn hợp đồng cần xóa!")
                return
            values = tree_hd.item(selected, 'values')
            mahd = values[0]
            confirm = messagebox.askyesno("Xác nhận xóa", f"Bạn có chắc chắn muốn xóa hợp đồng '{mahd}' không?")
            if not confirm:
                return
            try:
                cursor.execute("DELETE FROM HopDong WHERE MaHD=?", (mahd,))
                conn.commit()
                messagebox.showinfo("Thành công", f"Đã xóa hợp đồng '{mahd}'")
                load_data_hopdong()
            except Exception as e:
                messagebox.showerror("Lỗi", f"Không thể xóa:\n{e}")
        Button(nut_frame, text="Xóa", bg="#0D47A1", fg="white",
                font=("Times New Roman", 13, "bold"), width=10, command=xoa_hopdong).pack(pady=5)
       
        # Làm mới form
        def lam_moi_form_hd():
            entry_mahd.delete(0, END)
            try:
                entry_ngaybd.set_date(datetime.now())
                entry_ngaykt.set_date(datetime.now())
            except:
                pass
            entry_tengv_hd.delete(0, END)
            entry_mahd.focus()
        Button(nut_frame, text="Làm mới", bg="#0D47A1", fg="white",
                font=("Times New Roman", 13, "bold"), width=10, command=lam_moi_form_hd).pack(pady=5)
        
        #Lưu
        def luu_hopdong():
            if not tree_hd.get_children():
                messagebox.showwarning("Chưa có dữ liệu", "Danh sách hợp đồng trống!")
                return
            wb = Workbook()
            ws = wb.active
            ws.title = "Danh sách hợp đồng"
            columns = ("Mã hợp đồng", "Ngày bắt đầu", "Ngày kết thúc", "Tên giáo viên")
            ws.append(columns)
            for row_id in tree_hd.get_children():
                values = tree_hd.item(row_id)['values']
                ws.append(values)
            file_path = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                                    filetypes=[("Excel files", "*.xlsx")])
            if file_path:
                try:
                    wb.save(file_path)
                    messagebox.showinfo("Thành công", f"Đã lưu file Excel: {file_path}")
                except Exception as e:
                    messagebox.showerror("Lỗi", f"Không thể lưu file:\n{e}")
        Button(nut_frame, text="Lưu", bg="#0D47A1", fg="white",
                font=("Times New Roman", 13, "bold"), width=10,command=lambda: luu_hopdong()).pack(pady=5)

        # Treeview
        table_frame = Frame(noi_dung_frame, bg="white")
        table_frame.pack(fill=BOTH, expand=True, padx=20, pady=10)

        columns = ("mahd", "ngaybd", "ngaykt", "tengv")
        tree_hd = ttk.Treeview(table_frame, columns=columns, show="headings", height=25)
        tree_hd.heading("mahd", text="Mã hợp đồng")
        tree_hd.heading("ngaybd", text="Ngày bắt đầu")
        tree_hd.heading("ngaykt", text="Ngày kết thúc")
        tree_hd.heading("tengv", text="Tên giáo viên")
        for col in columns:
            tree_hd.column(col, width=120, anchor=W)
        tree_hd.pack(fill=BOTH, expand=True)
        tree_hd.bind("<ButtonRelease-1>", hien_thi_entry_hd)

        load_data_hopdong()


    #=====QUẢN LÝ NGHỈ PHÉP=====
    def load_data_nghiphep():
        global tree
        for row in tree.get_children():
            tree.delete(row)
        try:
            cursor.execute("""
                SELECT n.MaNghiPhep, n.NgayNghi, n.SoNgayNghi, n.LyDo, g.HoTenGV
                FROM NghiPhep n
                JOIN GiaoVien g ON n.MaGV = g.MaGV
            """)
            rows = cursor.fetchall()
            for row in rows:
                tree.insert("", "end", values=[str(col) for col in row])
        except Exception as e:
            messagebox.showerror("Lỗi truy vấn", f"Không thể tải dữ liệu nghỉ phép:\n{e}")

    def hien_thi_nghi_phep():
        global tree
        for widget in noi_dung_frame.winfo_children():
            widget.destroy()

        Label(noi_dung_frame, text="QUẢN LÝ NGHỈ PHÉP",
            font=("Times New Roman", 24, "bold"),
            fg="#0a46a3", bg="white").pack(pady=10)

        khung_nhap = LabelFrame(noi_dung_frame, text="Nhập thông tin",
                                font=("Times New Roman", 14, "bold"),
                                bg="#cce4f7", fg="#0a46a3", padx=15, pady=10)
        khung_nhap.pack(fill=X, padx=20, pady=10)

        # Tìm kiếm
        search_frame = Frame(khung_nhap, bg="#cce4f7")
        search_frame.grid(row=0, column=0, columnspan=5, sticky=W, pady=5)

        global entry_timkiem
        Label(search_frame, text="Tìm kiếm:", bg="#cce4f7", 
              font=("Times New Roman", 13)).pack(side=LEFT, padx=5)
        entry_timkiem = Entry(search_frame, width=30, font=("Times New Roman", 13))
        entry_timkiem.pack(side=LEFT, padx=5)
        placeholder = "Nhập mã đơn nghỉ phép..."
        entry_timkiem.insert(0, placeholder)
        entry_timkiem.config(fg="grey")
        def on_focus_in(event):
            if entry_timkiem.get() == placeholder:
                entry_timkiem.delete(0, END)
                entry_timkiem.config(fg="black")
        def on_focus_out(event):
            if entry_timkiem.get().strip() == "":
                entry_timkiem.insert(0, placeholder)
                entry_timkiem.config(fg="grey")
        entry_timkiem.bind("<FocusIn>", on_focus_in)
        entry_timkiem.bind("<FocusOut>", on_focus_out)

        # Entry nhập liệu
        global entry_manghiphep, entry_ngaynghi, entry_songaynghi, entry_lydo, entry_hoten

        Label(khung_nhap, text="Mã đơn nghỉ phép:", bg="#cce4f7",
               font=("Times New Roman", 13)).grid(row=1, column=0, sticky=W, padx=5, pady=5)
        entry_manghiphep = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_manghiphep.grid(row=1, column=1, padx=5, pady=5)

        Label(khung_nhap, text="Ngày nghỉ:", bg="#cce4f7", 
              font=("Times New Roman", 13)).grid(row=1, column=2, sticky=W, padx=5, pady=5)
        entry_ngaynghi = DateEntry(khung_nhap, width=23, 
            font=("Times New Roman", 13), background='darkblue',
            foreground='white', borderwidth=2, date_pattern='yyyy-mm-dd')
        entry_ngaynghi.grid(row=1, column=3, padx=5, pady=5)

        Label(khung_nhap, text="Số ngày nghỉ:", bg="#cce4f7",
               font=("Times New Roman", 13)).grid(row=2, column=0, sticky=W, padx=5, pady=5)
        entry_songaynghi = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_songaynghi.grid(row=2, column=1, padx=5, pady=5)

        Label(khung_nhap, text="Lý do:", bg="#cce4f7",
               font=("Times New Roman", 13)).grid(row=2, column=2, sticky=W, padx=5, pady=5)
        entry_lydo = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_lydo.grid(row=2, column=3, padx=5, pady=5)

        Label(khung_nhap, text="Họ tên giáo viên:", bg="#cce4f7",
               font=("Times New Roman", 13)).grid(row=3, column=0, sticky=W, padx=5, pady=5)
        entry_hoten = Entry(khung_nhap, width=25, font=("Times New Roman", 13))
        entry_hoten.grid(row=3, column=1, padx=5, pady=5)

        # Nút CRUD
        nut_frame = Frame(khung_nhap, bg="#cce4f7")
        nut_frame.grid(row=1, column=4, rowspan=5, padx=20, sticky=N)

        def hien_thi_entry_np(event):
            selected = tree.focus()
            if not selected:
                return
            values = tree.item(selected, 'values')
            entry_manghiphep.delete(0, END)
            entry_manghiphep.insert(0, values[0])
            try:
                dt = datetime.strptime(values[1], "%Y-%m-%d")
                entry_ngaynghi.set_date(dt)
            except:
                pass
            entry_songaynghi.delete(0, END)
            entry_songaynghi.insert(0, values[2])
            entry_lydo.delete(0, END)
            entry_lydo.insert(0, values[3])
            entry_hoten.delete(0, END)
            entry_hoten.insert(0, values[4])

        # Thêm
        def them_nghiphep():
            manp = entry_manghiphep.get().strip()
            ngay = entry_ngaynghi.get_date().strftime("%Y-%m-%d")
            songay = entry_songaynghi.get().strip()
            lydo = entry_lydo.get().strip()
            hoten = entry_hoten.get().strip()

            if not manp or not hoten or not songay:
                messagebox.showwarning("Thiếu dữ liệu", 
                                       "Vui lòng nhập đầy đủ các trường bắt buộc!")
                return
            try:
                # Lấy MaGV từ họ tên
                cursor.execute("SELECT MaGV FROM GiaoVien WHERE HoTenGV=?", (hoten,))
                row = cursor.fetchone()
                if not row:
                    messagebox.showerror("Lỗi", f"Không tìm thấy giáo viên '{hoten}'")
                    return
                magv = row[0]
                # Kiểm tra tồn tại MaNghiPhep
                cursor.execute("SELECT * FROM NghiPhep WHERE MaNghiPhep=?", (manp,))
                if cursor.fetchone():
                    messagebox.showwarning("Bản ghi tồn tại", f"Mã đơn '{manp}' đã tồn tại!")
                    return
                cursor.execute("INSERT INTO NghiPhep(MaNghiPhep, NgayNghi, SoNgayNghi, LyDo, MaGV) VALUES(?,?,?,?,?)",
                            (manp, ngay, songay, lydo if lydo else None, magv))
                conn.commit()
                messagebox.showinfo("Thành công", "Đã thêm đơn nghỉ phép mới!")
                load_data_nghiphep()
            except Exception as e:
                messagebox.showerror("Lỗi", f"Không thể thêm:\n{e}")

        Button(nut_frame, text="Thêm", bg="#0D47A1", fg="white",
            font=("Times New Roman", 13, "bold"), width=10, command=them_nghiphep).pack(pady=5)

        # Sửa
        def sua_nghiphep():
            manp = entry_manghiphep.get().strip()
            ngay = entry_ngaynghi.get_date().strftime("%Y-%m-%d")
            songay = entry_songaynghi.get().strip()
            lydo = entry_lydo.get().strip()
            hoten = entry_hoten.get().strip()

            if not manp:
                messagebox.showwarning("Thiếu dữ liệu", "Vui lòng chọn một bản ghi để sửa!")
                return
            try:
                cursor.execute("SELECT * FROM NghiPhep WHERE MaNghiPhep=?", (manp,))
                if not cursor.fetchone():
                    messagebox.showwarning("Không tồn tại", f"Mã đơn '{manp}' không tồn tại!")
                    return
                # Lấy MaGV từ họ tên
                cursor.execute("SELECT MaGV FROM GiaoVien WHERE HoTenGV=?", (hoten,))
                row = cursor.fetchone()
                if not row:
                    messagebox.showerror("Lỗi", f"Không tìm thấy giáo viên '{hoten}'")
                    return
                magv = row[0]
                cursor.execute("UPDATE NghiPhep SET NgayNghi=?, SoNgayNghi=?, LyDo=?, MaGV=? WHERE MaNghiPhep=?",
                            (ngay, songay, lydo if lydo else None, magv, manp))
                conn.commit()
                messagebox.showinfo("Thành công", "Đã sửa đơn nghỉ phép!")
                load_data_nghiphep()
            except Exception as e:
                messagebox.showerror("Lỗi", f"Không thể sửa:\n{e}")

        Button(nut_frame, text="Sửa", bg="#0D47A1", fg="white",
            font=("Times New Roman", 13, "bold"), width=10, command=sua_nghiphep).pack(pady=5)

        # Xóa
        def xoa_nghiphep():
            selected = tree.focus()
            if not selected:
                messagebox.showwarning("Chưa chọn", "Vui lòng chọn đơn cần xóa!")
                return
            values = tree.item(selected, 'values')
            manp = values[0]
            confirm = messagebox.askyesno("Xác nhận xóa", f"Bạn có chắc chắn muốn xóa đơn '{manp}' không?")
            if not confirm:
                return
            try:
                cursor.execute("DELETE FROM NghiPhep WHERE MaNghiPhep=?", (manp,))
                conn.commit()
                messagebox.showinfo("Thành công", f"Đã xóa đơn '{manp}'")
                load_data_nghiphep()
            except Exception as e:
                messagebox.showerror("Lỗi", f"Không thể xóa:\n{e}")

        Button(nut_frame, text="Xóa", bg="#0D47A1", fg="white",
            font=("Times New Roman", 13, "bold"), width=10, command=xoa_nghiphep).pack(pady=5)

        # Làm mới form
        def lam_moi_form_np():
            entry_manghiphep.delete(0, END)
            try:
                entry_ngaynghi.set_date(datetime.now())
            except:
                pass
            entry_songaynghi.delete(0, END)
            entry_lydo.delete(0, END)
            entry_hoten.delete(0, END)
            entry_manghiphep.focus()

        Button(nut_frame, text="Làm mới", bg="#0D47A1", fg="white",
               font=("Times New Roman", 13, "bold"), width=10, command=lam_moi_form_np).pack(pady=5)
        
        #Lưu
        def luu_nghiphep():
            if not tree.get_children():
                messagebox.showwarning("Chưa có dữ liệu", "Danh sách nghỉ phép trống!")
                return
            wb = Workbook()
            ws = wb.active
            ws.title = "Danh sách nghỉ phép"
            # Tiêu đề cột
            columns = ("Mã đơn nghỉ phép", "Ngày nghỉ", "Số ngày nghỉ", "Lý do", "Tên giáo viên")
            ws.append(columns)
            # Thêm dữ liệu từ Treeview
            for row_id in tree.get_children():
                values = tree.item(row_id)['values']
                ws.append(values)
            # Chọn đường dẫn lưu file
            file_path = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                                    filetypes=[("Excel files", "*.xlsx")])
            if file_path:
                try:
                    wb.save(file_path)
                    messagebox.showinfo("Thành công", f"Đã lưu file Excel: {file_path}")
                except Exception as e:
                    messagebox.showerror("Lỗi", f"Không thể lưu file:\n{e}")
        Button(nut_frame, text="Lưu", bg="#0D47A1", fg="white",
               font=("Times New Roman", 13, "bold"), width=10,command=luu_nghiphep).pack(pady=5)

        # Tìm kiếm theo Mã nghỉ phép
        def tim_nghiphep():
            manp_tim = entry_timkiem.get().strip()
            if manp_tim == "" or manp_tim == placeholder:
                load_data_nghiphep()
                return
            for row in tree.get_children():
                tree.delete(row)
            try:
                cursor.execute("""
                    SELECT n.MaNghiPhep, n.NgayNghi, n.SoNgayNghi, n.LyDo, g.HoTenGV
                    FROM NghiPhep n
                    JOIN GiaoVien g ON n.MaGV = g.MaGV
                    WHERE n.MaNghiPhep=?
                """, (manp_tim,))
                rows = cursor.fetchall()
                if rows:
                    for row in rows:
                        tree.insert("", "end", values=[str(col) for col in row])
                else:
                    messagebox.showinfo("Không tìm thấy", f"Không tìm thấy đơn '{manp_tim}'")
                    load_data_nghiphep()
            except Exception as e:
                messagebox.showerror("Lỗi", f"Không thể tìm:\n{e}")

        Button(search_frame, text="Tìm", bg="#e80101", fg="white", font=("Times New Roman", 13, "bold"),
            command=tim_nghiphep).pack(side=LEFT, padx=5)

        # Treeview
        table_frame = Frame(noi_dung_frame, bg="white")
        table_frame.pack(fill=BOTH, expand=True, padx=20, pady=10)

        columns = ("manghiphep", "ngaynghi", "songaynghi", "lydo", "hoten")
        tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=25)

        tree.heading("manghiphep", text="Mã đơn nghỉ phép")
        tree.heading("ngaynghi", text="Ngày nghỉ")
        tree.heading("songaynghi", text="Số ngày nghỉ")
        tree.heading("lydo", text="Lý do")
        tree.heading("hoten", text="Họ tên giáo viên")

        for col in columns:
            tree.column(col, width=120, anchor=W)

        tree.pack(fill=BOTH, expand=True)
        tree.bind("<ButtonRelease-1>", hien_thi_entry_np)

        load_data_nghiphep()


    #DÒNG CHÀO MỪNG
    Label(main,
          text="CHÀO MỪNG BẠN ĐẾN VỚI CHƯƠNG TRÌNH\nQUẢN LÝ GIÁO VIÊN THPT",
          fg="#0a46a3",
          bg="white",
          justify="center",
          font=("Times New Roman", 30, "bold")
          ).pack(expand=True, fill=BOTH)

    #CHÂN TRANG
    Label(main, text="Phiên bản 1.0 - © 2025", bg="white", fg="gray",
          font=("Times New Roman", 10, "italic")).pack(side=BOTTOM, pady=10)

 #=====Load dữ liệu ban đầu=====   
root.mainloop()